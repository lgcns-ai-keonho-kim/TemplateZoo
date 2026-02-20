"""
목적: ingestion 입력 파일을 스캔하고 텍스트를 추출한다.
설명: 확장자 필터링, 파일별 텍스트 추출, 헤더 구조 추정 메타데이터 생성을 담당한다.
디자인 패턴: 함수형 파서 모듈
참조: ingestion/core/chunking.py, ingestion/steps/chunk_step.py
"""

from __future__ import annotations

import re
from html import escape
from statistics import median
from pathlib import Path

from ingestion.core.docx_layout import (
    advance_docx_page_state as _advance_docx_page_state,
    estimate_docx_paragraph_layout as _estimate_docx_paragraph_layout,
    estimate_docx_table_height as _estimate_docx_table_height,
    is_docx_page_break_before as _is_docx_page_break_before,
    iterate_docx_blocks as _iterate_docx_blocks,
    resolve_docx_layout_metrics as _resolve_docx_layout_metrics,
)
from ingestion.core.image_annotation import ImageAnnotationTask, annotate_image_tasks
from ingestion.core.pdf_assets import PdfTextBlock, extract_pdf_page_assets
from ingestion.core.pdf_tables import extract_pdf_tables_by_page
from ingestion.core.table_annotation import TableAnnotationTask, annotate_table_tasks
from rag_chatbot.shared.exceptions import BaseAppException, ExceptionDetail

SUPPORTED_SUFFIXES = {".pdf", ".md", ".markdown", ".docx"}


def scan_input_files(input_root: str | Path) -> list[Path]:
    """ingestion 입력 디렉터리를 스캔한다."""

    root = Path(input_root)
    if not root.exists():
        detail = ExceptionDetail(
            code="INGESTION_INPUT_NOT_FOUND",
            cause=f"input_root={root}",
        )
        raise BaseAppException("ingestion 입력 경로를 찾을 수 없습니다.", detail)

    return [
        path
        for path in sorted(root.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    ]


def extract_text_by_file(path: Path) -> list[tuple[str, dict[str, object]]]:
    """파일에서 텍스트와 레이아웃 메타데이터를 추출한다."""

    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        text = path.read_text(encoding="utf-8")
        return [(text, {"layout_type": "markdown", "page_num": None})]
    if suffix == ".docx":
        return _read_docx_blocks(path)
    if suffix == ".pdf":
        return _read_pdf_blocks(path)
    detail = ExceptionDetail(
        code="INGESTION_FILE_TYPE_UNSUPPORTED",
        cause=f"path={path}",
    )
    raise BaseAppException("지원하지 않는 파일 형식입니다.", detail)


def _read_docx_blocks(path: Path) -> list[tuple[str, dict[str, object]]]:
    try:
        from docx import Document as WordDocument
    except ImportError as error:
        detail = ExceptionDetail(
            code="INGESTION_DOCX_READER_MISSING",
            cause="python-docx 패키지가 설치되어 있지 않습니다.",
        )
        raise BaseAppException("DOCX ingestion을 위해 python-docx가 필요합니다.", detail, error) from error

    document = WordDocument(str(path))
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
    body_font_size = _estimate_docx_body_font_size(paragraphs)
    layout_metrics = _resolve_docx_layout_metrics(document)
    usable_width_pt = float(layout_metrics["usable_width_pt"])
    usable_height_pt = float(layout_metrics["usable_height_pt"])

    rows: list[tuple[str, dict[str, object]]] = []
    page_texts: dict[int, list[str]] = {}
    table_placeholders: list[tuple[str, int, str]] = []
    table_placeholder_prefix = "__TABLE_TASK__::"
    paragraph_index = 0
    table_index = 0
    current_page_num = 1
    used_height_pt = 0.0

    for block_type, block in _iterate_docx_blocks(document):
        if block_type == "paragraph":
            paragraph = block
            text = str(paragraph.text or "").strip()
            if not text:
                continue

            if _is_docx_page_break_before(paragraph):
                current_page_num += 1
                used_height_pt = 0.0

            paragraph_index += 1
            font_size = _max_docx_font_size(paragraph) or body_font_size
            style_name = ""
            if paragraph.style is not None and paragraph.style.name is not None:
                style_name = str(paragraph.style.name).strip()
            heading_level = _resolve_docx_heading_level(
                style_name=style_name,
                font_size=font_size,
                body_font_size=body_font_size,
                text=text,
            )
            paragraph_layout = _estimate_docx_paragraph_layout(
                paragraph=paragraph,
                text=text,
                font_size=font_size,
                usable_width_pt=usable_width_pt,
            )
            block_height_pt = float(paragraph_layout["estimated_height_pt"])
            page_num, current_page_num, used_height_pt = _advance_docx_page_state(
                current_page_num=current_page_num,
                used_height_pt=used_height_pt,
                block_height_pt=block_height_pt,
                usable_height_pt=usable_height_pt,
            )

            metadata: dict[str, object] = {
                "layout_type": "docx_paragraph",
                "page_num": page_num,
                "block_type": "heading" if heading_level is not None else "paragraph",
                "heading_tag": f"H{heading_level}" if heading_level is not None else "BODY",
                "paragraph_index": paragraph_index,
                "font_size": round(float(font_size), 2),
                "body_font_size": round(float(body_font_size), 2),
                "line_spacing": round(float(paragraph_layout["line_spacing_pt"]), 2),
                "char_spacing": round(float(paragraph_layout["char_spacing_pt"]), 2),
                "space_before": round(float(paragraph_layout["space_before_pt"]), 2),
                "space_after": round(float(paragraph_layout["space_after_pt"]), 2),
                "line_count_estimated": int(paragraph_layout["line_count_estimated"]),
                "estimated_height_pt": round(block_height_pt, 2),
                "page_size": "A4",
            }
            if heading_level is not None:
                metadata["heading_level"] = heading_level
            rows.append((text, metadata))
            page_texts.setdefault(page_num, []).append(text)
            continue

        table = block
        table_html = _serialize_docx_table(table)
        if not table_html:
            continue
        table_index += 1
        table_height_pt = _estimate_docx_table_height(
            table=table,
            body_font_size=body_font_size,
            usable_width_pt=usable_width_pt,
        )
        page_num, current_page_num, used_height_pt = _advance_docx_page_state(
            current_page_num=current_page_num,
            used_height_pt=used_height_pt,
            block_height_pt=table_height_pt,
            usable_height_pt=usable_height_pt,
        )
        task_id = f"docx::{path.name}::{table_index}"
        table_placeholders.append((task_id, page_num, table_html))
        rows.append(
            (
                f"{table_placeholder_prefix}{task_id}",
                {
                    "layout_type": "docx_table",
                    "page_num": page_num,
                    "block_type": "table",
                    "table_index": table_index,
                    "heading_tag": "BODY",
                    "body_font_size": round(float(body_font_size), 2),
                    "estimated_height_pt": round(float(table_height_pt), 2),
                    "page_size": "A4",
                },
            )
        )

    table_tasks: list[TableAnnotationTask] = []
    for task_id, page_num, table_html in table_placeholders:
        page_text = "\n".join(page_texts.get(page_num, [])).strip()
        table_tasks.append(
            TableAnnotationTask(
                task_id=task_id,
                page_num=page_num,
                table_html=table_html,
                page_text=page_text,
            )
        )
    table_bodies = annotate_table_tasks(table_tasks)

    resolved_rows: list[tuple[str, dict[str, object]]] = []
    for row_text, row_metadata in rows:
        if row_text.startswith(table_placeholder_prefix):
            task_id = row_text.removeprefix(table_placeholder_prefix)
            table_body = table_bodies.get(task_id)
            if table_body is None:
                detail = ExceptionDetail(
                    code="INGESTION_TABLE_ANNOTATION_MISSING",
                    cause=f"task_id={task_id}",
                )
                raise BaseAppException("표 주석 결과를 찾을 수 없습니다.", detail)
            resolved_rows.append((table_body, row_metadata))
            continue
        resolved_rows.append((row_text, row_metadata))
    return resolved_rows


def _read_pdf_blocks(path: Path) -> list[tuple[str, dict[str, object]]]:
    page_assets = extract_pdf_page_assets(path)
    table_htmls_by_page = extract_pdf_tables_by_page(path)

    page_text_lookup = {page.page_num: page.page_text for page in page_assets}
    table_tasks: list[TableAnnotationTask] = []
    for page_num, table_htmls in table_htmls_by_page.items():
        page_text = page_text_lookup.get(page_num, "")
        for table_index, table_html in enumerate(table_htmls, start=1):
            task_id = f"pdf::{path.name}::p{page_num}::t{table_index}"
            table_tasks.append(
                TableAnnotationTask(
                    task_id=task_id,
                    page_num=page_num,
                    table_html=table_html,
                    page_text=page_text,
                )
            )
    table_bodies = annotate_table_tasks(table_tasks)

    table_rows_by_page: dict[int, list[tuple[str, dict[str, object]]]] = {}
    for task in table_tasks:
        table_body = table_bodies.get(task.task_id)
        if table_body is None:
            detail = ExceptionDetail(
                code="INGESTION_TABLE_ANNOTATION_MISSING",
                cause=f"task_id={task.task_id}",
            )
            raise BaseAppException("표 주석 결과를 찾을 수 없습니다.", detail)
        table_rows_by_page.setdefault(task.page_num, []).append(
            (
                table_body,
                {
                    "layout_type": "pdf_table",
                    "page_num": task.page_num,
                    "block_type": "table",
                },
            )
        )

    image_tasks: list[ImageAnnotationTask] = []
    for page_asset in page_assets:
        for image_path in page_asset.image_paths:
            image_tasks.append(
                ImageAnnotationTask(
                    page_num=page_asset.page_num,
                    image_path=image_path,
                    page_text=page_asset.page_text,
                )
            )

    image_rows_by_page: dict[int, list[tuple[str, dict[str, object]]]] = {}
    for image_body, image_metadata in annotate_image_tasks(image_tasks):
        page_num = int(image_metadata.get("page_num", 1))
        image_rows_by_page.setdefault(page_num, []).append((image_body, image_metadata))

    rows: list[tuple[str, dict[str, object]]] = []
    for page_asset in page_assets:
        body_font_size = _estimate_pdf_body_font_size(page_asset.text_blocks)
        for block_index, block in enumerate(page_asset.text_blocks, start=1):
            text = block.text
            font_size = block.font_size
            heading_level = _infer_heading_level(
                font_size=font_size,
                body_font_size=body_font_size,
                text=text,
            )
            metadata: dict[str, object] = {
                "layout_type": "pdf_block",
                "page_num": page_asset.page_num,
                "block_type": "heading" if heading_level is not None else "paragraph",
                "heading_tag": f"H{heading_level}" if heading_level is not None else "BODY",
                "block_index": block_index,
                "body_font_size": round(float(body_font_size), 2),
            }
            if font_size is not None:
                metadata["font_size"] = round(float(font_size), 2)
            if heading_level is not None:
                metadata["heading_level"] = heading_level
            rows.append((text, metadata))

        rows.extend(table_rows_by_page.get(page_asset.page_num, []))
        rows.extend(image_rows_by_page.get(page_asset.page_num, []))

    return rows


def _estimate_pdf_body_font_size(blocks: list[PdfTextBlock]) -> float:
    sizes: list[float] = []
    for block in blocks:
        size = block.font_size
        if isinstance(size, int | float) and size > 0:
            sizes.append(float(size))
    return _estimate_body_font_size_from_samples(sizes)


def _estimate_docx_body_font_size(paragraphs: list[object]) -> float:
    sizes: list[float] = []
    for paragraph in paragraphs:
        size = _max_docx_font_size(paragraph)
        if size is not None and size > 0:
            sizes.append(float(size))
    return _estimate_body_font_size_from_samples(sizes)


def _estimate_body_font_size_from_samples(sizes: list[float]) -> float:
    if not sizes:
        return 11.0
    ordered = sorted(sizes)
    # 제목 폰트의 영향을 줄이기 위해 하위 60% 구간의 중앙값을 본문 기준으로 사용한다.
    cutoff = max(1, int(len(ordered) * 0.6))
    return float(median(ordered[:cutoff]))


def _max_docx_font_size(paragraph: object) -> float | None:
    sizes: list[float] = []
    for run in paragraph.runs:
        size = getattr(run.font, "size", None)
        if size is None:
            continue
        point = getattr(size, "pt", None)
        if point is None:
            continue
        sizes.append(float(point))
    if not sizes:
        return None
    return max(sizes)


def _resolve_docx_heading_level(
    *,
    style_name: str,
    font_size: float | None,
    body_font_size: float,
    text: str,
) -> int | None:
    style_level = _parse_heading_level(style_name)
    if style_level is not None:
        return style_level
    return _infer_heading_level(font_size=font_size, body_font_size=body_font_size, text=text)


def _parse_heading_level(style_name: str) -> int | None:
    if not style_name:
        return None
    style_lower = style_name.lower()
    matched = re.search(r"(heading|제목)\s*(\d+)", style_lower)
    if not matched:
        return None
    level = int(matched.group(2))
    return max(1, min(level, 6))


def _infer_heading_level(
    *,
    font_size: float | None,
    body_font_size: float,
    text: str,
) -> int | None:
    if font_size is None or body_font_size <= 0:
        return None
    normalized_text = " ".join(text.split())
    if not normalized_text:
        return None
    # 과도하게 긴 문장은 제목으로 간주하지 않는다.
    if len(normalized_text) > 120:
        return None

    ratio = float(font_size) / float(body_font_size)
    if ratio >= 1.65:
        return 1
    if ratio >= 1.45:
        return 2
    if ratio >= 1.30:
        return 3
    return None


def _serialize_docx_table(table: object) -> str:
    rows_html: list[str] = []
    for row in table.rows:
        cells_html: list[str] = []
        has_non_empty_cell = False
        for cell in row.cells:
            value = str(cell.text or "")
            normalized = value.strip()
            if normalized:
                has_non_empty_cell = True
            escaped_lines = [escape(line.strip()) for line in value.splitlines() if line.strip()]
            if escaped_lines:
                cell_value = "<br/>".join(escaped_lines)
            else:
                cell_value = ""
            cells_html.append(f"<td>{cell_value}</td>")
        if not cells_html:
            continue
        if not has_non_empty_cell:
            continue
        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
    if not rows_html:
        return ""
    return f"<table><tbody>{''.join(rows_html)}</tbody></table>"


__all__ = ["SUPPORTED_SUFFIXES", "scan_input_files", "extract_text_by_file"]
